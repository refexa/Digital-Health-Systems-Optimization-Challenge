from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "refexa_Digital Health System_Challenge.docx"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=(255, 255, 255))
        shade(table.rows[0].cells[index], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
            if len(table.rows) % 2 == 0:
                shade(cells[index], "EAF2F8")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_flow(document, title, steps):
    document.add_paragraph(title, style="Heading 3")
    table = document.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, step in enumerate(steps):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, step, bold=True)
        shade(cell, "D9EAF7" if index % 2 == 0 else "F4F7F9")
        if index < len(steps) - 1:
            cell.add_paragraph("      ->")
    document.add_paragraph("Figure: " + title + ". Arrows show the primary process direction.", style="Caption")


def add_bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 2"].font.name = "Aptos Display"
    styles["Heading 2"].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 3"].font.name = "Aptos Display"
    styles["Heading 3"].font.color.rgb = RGBColor(68, 68, 68)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Optimizing Outpatient Appointment Scheduling")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("A practical digital workflow redesign for better access, capacity use, and staff decisions")
    run.italic = True
    run.font.size = Pt(12)
    document.add_paragraph()
    add_table(document, ["Submission", "Details"], [
        ("Candidate", "Refexa"),
        ("Workflow studied", "Outpatient appointment booking and capacity management"),
        ("Prototype", "Flask and SQLite scheduling service with web interface"),
        ("Submission format", "Digital Health System Challenge document"),
    ], [1.5, 5.8])

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "Outpatient appointment scheduling is a high-volume coordination task that directly affects access, waiting time, staff workload, and the use of clinical capacity. "
        "This submission analyzes a simple current-state process in which patients or staff search for appointments and manually respond when a preferred slot is unavailable. "
        "It proposes a digital scheduling workflow that exposes real-time capacity, recommends the nearest available alternative, supports controlled priority booking, and produces operational metrics."
    )
    document.add_paragraph(
        "The accompanying prototype models seven days of 30-minute appointment slots. It provides booking, priority booking, capacity management, schedule export, and metrics endpoints. "
        "The design is intentionally incremental: it can improve coordination without requiring a full replacement of the clinical record system."
    )

    document.add_heading("1. Healthcare Workflow Analysis", level=1)
    document.add_heading("1.1 Scope and stakeholders", level=2)
    add_table(document, ["Stakeholder", "Need or responsibility"], [
        ("Patient", "Find an acceptable appointment quickly, receive clear confirmation, and avoid unnecessary waiting."),
        ("Reception or contact-centre staff", "Book appointments accurately and resolve full-slot requests efficiently."),
        ("Clinician and care team", "Receive a predictable schedule and protect time for clinically urgent cases."),
        ("Operations manager", "Balance capacity, monitor utilization, and identify access problems."),
        ("Digital system", "Maintain a reliable source of slot availability and booking history."),
    ], [2.0, 5.3])

    add_flow(document, "Current-state patient journey", [
        "Request appointment", "Search available slots", "Select preferred slot", "Slot is full?", "Call or message staff", "Book or wait", "Attend visit"
    ])
    document.add_paragraph(
        "Figure 1. Current-state map. When the preferred slot is full, the workflow can branch into manual searching, repeated contact, or delayed care. "
        "The prototype makes this failure point explicit through a full-slot response."
    )

    document.add_heading("1.2 Bottlenecks and root causes", level=2)
    add_table(document, ["Bottleneck", "Root cause", "Operational and patient impact"], [
        ("Full preferred slots", "Demand is not matched to nearby open capacity at the point of booking.", "More contacts, longer waits, and avoidable abandonment."),
        ("Limited capacity visibility", "Availability is distributed across staff knowledge or disconnected tools.", "Slow decisions and underused appointment capacity."),
        ("Priority conflicts", "Urgent or high-priority requests are not consistently represented in the allocation rule.", "Clinical urgency may compete with first-come-first-served booking."),
        ("Manual reporting", "Utilization and booking distribution are calculated after the fact or not at all.", "Managers cannot quickly identify demand, capacity, or access problems."),
        ("Inconsistent booking data", "Patient, slot, time, and priority details may be entered in different places.", "Rework, errors, and weaker operational decisions."),
    ], [1.5, 2.5, 3.3])

    document.add_heading("1.3 Why the problem matters", level=2)
    add_bullet(document, "Patients experience uncertainty when a requested slot is unavailable and no immediate alternative is offered.")
    add_bullet(document, "Staff spend time searching across the schedule instead of resolving exceptions and supporting patients.")
    add_bullet(document, "Unused capacity can coexist with long waits when nearby open slots are not visible or are difficult to compare.")
    add_bullet(document, "Without a transparent allocation rule, priority decisions may be inconsistent and difficult to audit.")

    document.add_heading("2. Digital Workflow Redesign", level=1)
    document.add_heading("2.1 Future-state workflow", level=2)
    add_flow(document, "Future-state digital booking workflow", [
        "Request and preference", "Check live capacity", "Recommend nearest slot", "Confirm booking", "Send status", "Monitor metrics"
    ])
    document.add_paragraph(
        "Figure 2. Future-state map. The system responds to a full slot with the nearest available alternative instead of ending the interaction. "
        "Staff-only functions are protected by an administrator token and are separated from routine patient booking."
    )

    document.add_heading("2.2 Digital capabilities", level=2)
    add_table(document, ["Capability", "How it works", "Problem addressed"], [
        ("Live slot inventory", "Stores slot time, capacity, and booking records in SQLite and exposes them through the slots endpoint.", "Improves shared visibility and reduces duplicate searching."),
        ("Nearest-available recommendation", "When a slot is full, compares requested time with open slots and returns the closest available option.", "Reduces delay and patient effort after a failed booking."),
        ("Priority booking", "Authorized staff can assign a priority and preempt a lower-priority booking when the rule allows it.", "Creates a controlled route for urgent cases."),
        ("Capacity management", "Authorized staff can update the capacity of a slot as staffing or room availability changes.", "Aligns the schedule with real operational capacity."),
        ("Operational metrics", "Reports total slots, bookings, utilization, and booking distribution by slot.", "Supports data-informed daily management."),
        ("CSV export", "Exports slot, capacity, booking count, and booking detail for review.", "Enables audit, reporting, and offline analysis."),
    ], [1.6, 3.3, 2.4])

    document.add_heading("2.3 Controls and patient-centered design", level=2)
    add_bullet(document, "Keep patients in control: show the requested slot and alternative before final confirmation.")
    add_bullet(document, "Use clear status responses: booked, full with suggestion, unauthorized, or not found.")
    add_bullet(document, "Restrict priority changes, capacity changes, and exports to authorized staff.")
    add_bullet(document, "Retain an audit trail containing booking time, patient label, slot, and priority; production use would require stronger identity, access, and privacy controls.")
    add_bullet(document, "Design for accessibility and multiple channels, including staff-assisted booking for patients without digital access.")

    document.add_heading("2.4 Implementation sequence", level=2)
    add_table(document, ["Phase", "Action", "Success signal"], [
        ("1. Baseline", "Measure booking demand, response time, full-slot rate, waiting time, and unused capacity.", "A reliable baseline exists by clinic and appointment type."),
        ("2. Pilot", "Deploy live availability and nearest-slot suggestions for one clinic or appointment type.", "Fewer abandoned or manually escalated requests."),
        ("3. Govern", "Add role-based access, audit logging, notifications, and privacy controls.", "Priority and capacity changes are traceable."),
        ("4. Scale", "Integrate with the patient record, reminders, cancellation, and waitlist services.", "Benefits persist across clinics and channels."),
    ], [1.2, 4.0, 2.1])

    document.add_heading("3. Impact Assessment", level=1)
    document.add_heading("3.1 Expected outcomes and measures", level=2)
    add_table(document, ["Outcome", "Measure", "Target for pilot"], [
        ("Reduced waiting", "Median time from request to confirmed appointment; time to alternative after a full slot.", "10-20% reduction after workflow stabilization."),
        ("Faster booking", "Average staff handling time and percentage of requests resolved in one interaction.", "15% faster handling; fewer repeat contacts."),
        ("Better capacity use", "Booked slots divided by available capacity, plus unused-slot rate by day.", "5-10 percentage-point improvement where spare capacity exists."),
        ("Better access", "Percentage of full-slot requests receiving an available alternative.", "At least 90% receive a recommendation when capacity exists."),
        ("Decision quality", "Percentage of priority changes with authorized user and recorded reason.", "100% of priority changes traceable."),
        ("Patient experience", "Post-booking satisfaction and complaints about appointment access.", "Improved satisfaction and fewer access complaints."),
    ], [1.5, 3.5, 2.3])
    document.add_paragraph(
        "These are pilot targets, not measured claims. They should be tested against a baseline using the same clinic, appointment type, and observation period. "
        "The prototype supplies the basic booking and distribution data needed for an initial operational dashboard, but a production evaluation should also collect timestamps, cancellations, no-shows, and patient feedback."
    )

    document.add_heading("3.2 Prototype evidence", level=2)
    add_table(document, ["Prototype fact", "Relevance to the redesign"], [
        ("Seven days of slots are initialized.", "Provides a bounded planning horizon for access and capacity experiments."),
        ("Each day contains sixteen 30-minute slots.", "Creates a transparent appointment template that can be adapted by clinic."),
        ("A slot has explicit capacity and booking records.", "Allows availability to be computed rather than inferred manually."),
        ("Full bookings return a nearest available suggestion when one exists.", "Demonstrates the core recovery path for a failed booking."),
        ("Priority booking and capacity changes require an admin token.", "Demonstrates separation between patient booking and staff controls."),
        ("Metrics and CSV export endpoints are available.", "Supports monitoring, audit, and follow-up analysis."),
    ], [3.1, 4.2])

    document.add_heading("3.3 Assumptions, risks, and mitigations", level=2)
    add_table(document, ["Assumption or risk", "Mitigation"], [
        ("Projected improvements may not generalize from a simulation.", "Run a controlled pilot and compare with a documented baseline."),
        ("Priority preemption could inconvenience an existing patient.", "Require clinical policy, staff authorization, reason capture, and patient notification."),
        ("A simple token is not sufficient for production healthcare data.", "Use enterprise identity, role-based access, encryption, audit logs, and privacy review."),
        ("A nearest-time rule may not reflect clinical suitability.", "Add appointment type, clinician, location, accessibility, and patient preference constraints."),
        ("Capacity may change because of staffing or rooms.", "Give authorized operations staff a controlled capacity update workflow."),
        ("Patients may lack digital access.", "Maintain phone and staff-assisted channels using the same live schedule."),
    ], [3.2, 4.1])

    document.add_heading("Conclusion", level=1)
    document.add_paragraph(
        "The central opportunity is to treat appointment access as a coordinated digital workflow rather than a sequence of isolated booking actions. "
        "A shared slot inventory, immediate alternatives, governed priority rules, and simple operational metrics address the main failure points while preserving staff oversight and patient choice. "
        "The prototype provides a practical starting point; the next step is a measured pilot with stronger security, integration, and patient feedback."
    )

    document.add_heading("Appendix: Prototype endpoints", level=1)
    add_table(document, ["Endpoint", "Purpose"], [
        ("GET /slots", "View slots, capacities, and bookings."),
        ("POST /book", "Book an available slot or receive a nearest-slot suggestion."),
        ("POST /book_priority", "Authorized priority booking and controlled preemption."),
        ("POST /set_capacity", "Authorized capacity update."),
        ("GET /metrics", "View slot and booking distribution metrics."),
        ("GET /export", "Authorized CSV schedule export."),
    ], [2.2, 5.1])
    document.add_paragraph("Prepared for the Digital Health Systems Optimization Challenge. All impact values labeled as targets are proposed evaluation criteria, not observed clinical outcomes.", style="Caption")
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)